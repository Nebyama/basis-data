import json
import openpyxl
from pathlib import Path

path = Path(r'D:\Judika Basis Data TI ORESTE.xlsx')
wb = openpyxl.load_workbook(path, data_only=True)

def parse_bobot(ws):
    rows = list(ws.iter_rows(values_only=True))
    return [
        {
            'no': int(row[0]),
            'nama': row[1],
            'bobot': float(row[2]),
            'rank_bobot': int(row[3]),
        }
        for row in rows[3:8]
    ]


def parse_alternatif(ws):
    rows = list(ws.iter_rows(values_only=True))
    out = []
    for row in rows[3:43]:
        if not row[0]:
            continue
        out.append({
            'no': int(row[0]),
            'nama': row[1],
            'c1': int(row[2]),
            'c2': int(row[3]),
            'c3': int(row[4]),
            'c4': int(row[5]),
            'c5': int(row[6]),
        })
    return out


def parse_besson(ws):
    rows = list(ws.iter_rows(values_only=True))
    besson = {}
    def assign(idx, ccode):
        for row in rows[idx[0]:idx[1]]:
            if not row[idx[2]]:
                continue
            alt = row[idx[2]+1]
            if not alt:
                continue
            if alt not in besson:
                besson[alt] = {}
            besson[alt][ccode] = {
                'nilai': float(row[idx[2]+2]),
                'besson_rank': float(row[idx[2]+3]),
                'normalisasi': float(row[idx[2]+4]),
            }
    # C1/C2/C3 area
    assign((3, 43, 0), 'c1')
    assign((3, 43, 7), 'c2')
    assign((3, 43, 14), 'c3')
    # C4/C5 area
    assign((49, 89, 0), 'c4')
    assign((49, 89, 7), 'c5')
    return besson


def parse_normalisasi(ws):
    rows = list(ws.iter_rows(values_only=True))
    out = []
    for row in rows[2:42]:
        if not row[0]:
            continue
        out.append({
            'no': int(row[0]),
            'nama': row[1],
            'c1': float(row[2]),
            'c2': float(row[3]),
            'c3': float(row[4]),
            'c4': float(row[5]),
            'c5': float(row[6]),
        })
    return out


def parse_distance(ws):
    rows = list(ws.iter_rows(values_only=True))
    out = []
    for row in rows[3:43]:
        if not row[0]:
            continue
        out.append({
            'no': int(row[0]),
            'nama': row[1],
            'd_c1': float(row[2]),
            'd_c2': float(row[3]),
            'd_c3': float(row[4]),
            'd_c4': float(row[5]),
            'd_c5': float(row[6]),
            'akumulasi': float(row[7]),
        })
    return out


def parse_ranking(ws):
    rows = list(ws.iter_rows(values_only=True))
    out = []
    for row in rows[2:42]:
        if not row[0]:
            continue
        out.append({
            'ranking': int(row[0]),
            'nama': row[1],
            'akumulasi': float(row[2]),
            'keterangan': row[3] if row[3] else None,
        })
    return out


data = {
    'bobot': parse_bobot(wb['Bobot Kriteria']),
    'alternatif': parse_alternatif(wb['Nilai Alternatif']),
    'besson': parse_besson(wb['Besson Rank']),
    'normalisasi': parse_normalisasi(wb['Normalisasi']),
    'distance_score': parse_distance(wb['Distance Score']),
    'ranking': parse_ranking(wb['Perankingan Akhir']),
}

outfile = Path(r'D:\basis data\oreste_data.json')
outfile.write_text(json.dumps(data, indent=2, ensure_ascii=False))
print('Exported', outfile)
